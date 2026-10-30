"""T20 workplace: выгрузка досье в Word и PDF (S10; пункты 25, 30)."""
import io

from docx import Document
from factories import hh_item, profile, source

from hr import db, decisions
from hr.normalize import normalize
from hr.pipeline import process

SECTIONS = ["Сводка", "Сшивка", "Резюме", "Места работы", "Навыки", "Цифровой след", "Публичная активность",
            "Хронология", "Оценки", "Взаимодействие", "Доступ"]


def person(conn, clock):
    profile(conn)
    it = hh_item()
    pid = process(conn, source(conn, "hh"), normalize("hh", it), it, clock.now()).person_id
    uid = db.val(conn, "INSERT INTO app_user (login, full_name, role) VALUES ('t', 'Т', 'recruiter') RETURNING id")
    cid = decisions.take(conn, pid, uid, clock.now())
    db.run(conn, "INSERT INTO contact (candidate_id, kind, value, value_hash) VALUES (%s, 'email', 'x@y.ru', 'h')",
           (cid,))
    return pid


def test_s10_docx_has_all_sections_and_no_contacts(conn, api, clock):
    pid = person(conn, clock)
    c, h, uid = api()
    r = c.get(f"/api/person/{pid}/export", params={"f": "docx"}, headers=h)
    assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for s in SECTIONS:
        assert s in heads, s
    text = "\n".join(p.text for p in doc.paragraphs) + "".join(c.text for t in doc.tables for row in t.rows
                                                                for c in row.cells)
    assert "x@y.ru" not in text and "Altium Designer" in text
    exp = db.one(conn, "SELECT format, author_id, person_id FROM dossier_export ORDER BY id DESC LIMIT 1")
    assert exp == {"format": "docx", "author_id": uid, "person_id": pid}
    assert db.val(conn, "SELECT count(*) FROM audit_log WHERE action = 'export' AND object_id = %s", (pid,)) == 1


def test_s10_pdf_records_export(conn, api, clock):
    pid = person(conn, clock)
    c, h, uid = api()
    r = c.get(f"/api/person/{pid}/export", params={"f": "pdf"}, headers=h)
    assert r.status_code == 200 and r.json()["export_id"]
    assert db.val(conn, "SELECT format FROM dossier_export WHERE id = %s", (r.json()["export_id"],)) == "pdf"
    assert db.val(conn, "SELECT count(*) FROM dossier_export") == 1


def test_export_needs_csrf_and_permission(conn, api, clock):
    pid = person(conn, clock)
    c, h, _ = api()
    assert c.get(f"/api/person/{pid}/export", params={"f": "docx"}).status_code == 403, "GET выгрузки — с CSRF"
    m, mh, _ = api("hiring_manager")
    assert m.get(f"/api/person/{pid}/export", params={"f": "docx"}, headers=mh).status_code == 403
    assert c.get(f"/api/person/{pid}/export", params={"f": "xls"}, headers=h).status_code == 422
