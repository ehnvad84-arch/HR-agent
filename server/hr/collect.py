"""Сбор резюме: официальное API → выгрузка из кабинета. Сбора со страниц нет (раздел 7 ТЗ).

Ограничения зашиты в код: не чаще запроса в секунду на площадку, не больше
crawl_max_pages страниц за прогон; если площадка показывает проверку на робота,
задача останавливается, причина пишется в crawl_queue.last_error.
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import shutil
import threading
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterator

import httpx

from . import pipeline
from .config import get_config
from .db import j, one, rows, run, setting, set_setting, val
from .normalize import normalize

log = logging.getLogger(__name__)


class CrawlBlocked(RuntimeError):
    """Площадка просит пройти проверку — обходить её мы не будем."""


class NotConfigured(RuntimeError):
    pass


class RateLimiter:
    def __init__(self, per_sec: float = 1.0):
        self.interval = 1.0 / per_sec
        self.last = 0.0
        self.lock = threading.Lock()

    def wait(self) -> None:
        with self.lock:
            delta = time.monotonic() - self.last
            if delta < self.interval:
                time.sleep(self.interval - delta)
            self.last = time.monotonic()


_limiters: dict[str, RateLimiter] = {}


def limiter(code: str) -> RateLimiter:
    return _limiters.setdefault(code, RateLimiter(1.0))


def source(conn, code: str) -> dict:
    s = one(conn, "SELECT * FROM source WHERE code = %s", (code,))
    if not s:
        raise NotConfigured(f"источник {code} не заведён в таблице source")
    return s


# ================================================================== hh.ru

class HH:
    """Клиент API hh.ru для работодателя. Токены хранятся в settings.hh_tokens."""

    def __init__(self, conn, client: httpx.Client | None = None):
        self.cfg = get_config()
        self.conn = conn
        self.http = client or httpx.Client(timeout=30, headers={"User-Agent": self.cfg.hh_user_agent,
                                                                 "HH-User-Agent": self.cfg.hh_user_agent})

    # --- OAuth
    def authorize_url(self, redirect_uri: str) -> str:
        return (f"{self.cfg.hh_oauth}/oauth/authorize?response_type=code&client_id={self.cfg.hh_client_id}"
                f"&redirect_uri={redirect_uri}")

    def _store(self, data: dict) -> dict:
        tokens = {"access_token": data["access_token"], "refresh_token": data.get("refresh_token"),
                  "expires_at": (datetime.now(timezone.utc) + timedelta(seconds=int(data.get("expires_in", 1209600)))).isoformat()}
        set_setting(self.conn, "hh_tokens", tokens, None, "обновлён токен hh")
        return tokens

    def exchange_code(self, code: str, redirect_uri: str) -> dict:
        r = self.http.post(f"{self.cfg.hh_api}/token", data={
            "grant_type": "authorization_code", "client_id": self.cfg.hh_client_id,
            "client_secret": self.cfg.hh_client_secret, "code": code, "redirect_uri": redirect_uri})
        r.raise_for_status()
        return self._store(r.json())

    def token(self) -> str:
        t = setting(self.conn, "hh_tokens") or {}
        if t.get("access_token") and datetime.fromisoformat(t["expires_at"]) > datetime.now(timezone.utc) + timedelta(minutes=5):
            return t["access_token"]
        refresh = t.get("refresh_token") or self.cfg.hh_refresh_token
        if not (refresh and self.cfg.hh_client_id):
            raise NotConfigured("hh: нет refresh_token — пройдите /api/oauth/hh/start")
        r = self.http.post(f"{self.cfg.hh_api}/token", data={"grant_type": "refresh_token", "refresh_token": refresh})
        if r.status_code >= 400:
            raise NotConfigured(f"hh: токен не обновился ({r.status_code}): {r.text[:200]}")
        return self._store(r.json())["access_token"]

    # --- запросы
    def get(self, path: str, params: dict | None = None, auth: bool = True) -> dict:
        limiter("hh").wait()
        headers = {"Authorization": f"Bearer {self.token()}"} if auth else {}
        r = self.http.get(f"{self.cfg.hh_api}{path}", params=params, headers=headers)
        if r.status_code == 403 and "captcha" in r.text.lower():
            raise CrawlBlocked("hh просит пройти проверку на робота — обход остановлен")
        if r.status_code == 429:
            raise CrawlBlocked("hh: превышен лимит запросов — обход остановлен до следующего прогона")
        r.raise_for_status()
        return r.json()

    def search_params(self, profile: dict) -> dict:
        kws = [k for k in (profile.get("keywords") or [])[:6] if k]
        text = profile["title"] if not kws else f"{profile['title']} OR " + " OR ".join(f'"{k}"' if " " in k else k for k in kws)
        geo = profile.get("geo") or "krasnodar"
        params: dict[str, Any] = {"text": text, "per_page": 50, "order_by": "publication_time"}
        if geo == "krasnodar":
            params["area"] = str(setting(self.conn, "hh_area", 53))
        elif geo == "relocation_ru":
            params["area"] = str(setting(self.conn, "hh_area", 53))
            params["relocation"] = "living_or_relocation"
        else:
            params["area"] = str(setting(self.conn, "hh_area_relocation", 113))
            params["schedule"] = "remote"
        return params

    def search(self, profile: dict, max_pages: int) -> Iterator[dict]:
        params = self.search_params(profile)
        for page in range(max_pages):
            data = self.get("/resumes", {**params, "page": page})
            for item in data.get("items", []):
                yield item
            if page + 1 >= int(data.get("pages", 0)):
                break

    def resume(self, resume_id: str) -> dict:
        return self.get(f"/resumes/{resume_id}")

    def employer_vacancies(self) -> list[dict]:
        if not self.cfg.hh_employer_id:
            raise NotConfigured("HH_EMPLOYER_ID не задан")
        out, page = [], 0
        while True:
            data = self.get("/vacancies", {"employer_id": self.cfg.hh_employer_id, "per_page": 100, "page": page}, auth=False)
            out += data.get("items", [])
            page += 1
            if page >= int(data.get("pages", 0)):
                return out

    def responses(self, vacancy_external_id: str) -> list[dict]:
        data = self.get("/negotiations/response", {"vacancy_id": vacancy_external_id, "per_page": 50})
        return [i.get("resume") for i in data.get("items", []) if i.get("resume")]


def open_contact_hh(conn, external_id: str, client: httpx.Client | None = None) -> dict:
    """Открытие контакта — платное действие по подписке, только по нажатию рекрутера."""
    full = HH(conn, client).resume(external_id)
    fio = " ".join(filter(None, [full.get("last_name"), full.get("first_name"), full.get("middle_name")]))
    contacts = []
    for c in full.get("contact") or []:
        kind = (c.get("type") or {}).get("id")
        v = c.get("value")
        if isinstance(v, dict):
            v = v.get("formatted") or "".join(str(v.get(k, "")) for k in ("country", "city", "number"))
        if v:
            contacts.append({"kind": "phone" if kind == "cell" or kind == "home" or kind == "work" else kind, "value": v,
                             "is_primary": bool(c.get("preferred"))})
    return {"fio": fio or None, "birth_date": full.get("birth_date"), "contacts": contacts, "raw": full}


# ================================================================== Работа России

class Trudvsem:
    def __init__(self, conn, client: httpx.Client | None = None):
        self.cfg = get_config()
        self.conn = conn
        self.http = client or httpx.Client(timeout=30)

    def cvs(self, max_pages: int) -> Iterator[dict]:
        region = setting(self.conn, "trudvsem_region", "2300000000000")
        for page in range(max_pages):
            limiter("trudvsem").wait()
            r = self.http.get(f"{self.cfg.trudvsem_base.rstrip('/')}/cvs/region/{region}",
                              params={"offset": page, "limit": 100})
            if r.status_code == 404:
                break
            r.raise_for_status()
            data = r.json()
            res = data.get("results") or {}
            items = res.get("cvs") or res.get("cv") or []
            if not items:
                break
            yield from items


def _relevant(profiles: list[dict], title: str | None, skills: list[str]) -> bool:
    """У открытых данных нет поиска по тексту: берём только похожее на наши должности."""
    from .scoring import trigram_sim
    hay = " ".join(skills).lower() + " " + (title or "").lower()
    for p in profiles:
        if trigram_sim(title, p["title"]) >= 0.3:
            return True
        if sum(1 for k in p["keywords"] or [] if k.lower() in hay) >= 2:
            return True
    return False


# ================================================================== прогон

def _job(conn, source_id: int, profile_id: int | None, params: dict) -> int:
    return val(conn, """INSERT INTO crawl_queue (source_id, role_profile_id, params, state, attempts)
                        VALUES (%s, %s, %s, 'running', 1) RETURNING id""", (source_id, profile_id, j(params)))


def _finish(conn, job_id: int, ok: bool, error: str | None = None) -> None:
    run(conn, "UPDATE crawl_queue SET state = %s, last_error = %s, finished_at = now() WHERE id = %s",
        ("done" if ok else "error", error, job_id))


def ingest_items(conn_factory, src_code: str, items: Iterator[dict], stats: dict,
                 filter_fn=None) -> None:
    """Каждое резюме — своя транзакция: одно битое не роняет весь прогон."""
    for item in items:
        stats["seen"] += 1
        try:
            with conn_factory() as conn:
                src = source(conn, src_code)
                n = normalize(src_code, item)
                if filter_fn and not filter_fn(n):
                    stats["skipped"] += 1
                    continue
                r = pipeline.process(conn, src, n, item)
                stats["new"] += int(r["new_obs"])
                stats["changed"] += int(r["changed"])
                stats["persons"] += int(r["person_created"])
                stats["alerts"] += r["alerts"]
        except (CrawlBlocked, NotConfigured):
            raise
        except Exception as e:  # noqa: BLE001 — битая запись пишется в статистику, прогон идёт дальше
            log.exception("резюме не принято")
            stats["errors"] += 1
            stats.setdefault("last_errors", []).append(str(e)[:200])


def _stats() -> dict:
    return {"seen": 0, "new": 0, "changed": 0, "persons": 0, "alerts": 0, "skipped": 0, "errors": 0}


def collect_hh(conn_factory, client: httpx.Client | None = None) -> dict:
    stats = _stats()
    with conn_factory() as conn:
        src = source(conn, "hh")
        if not src["is_active"]:
            return {**stats, "note": "источник выключен"}
        profiles = pipeline.active_profiles(conn)
        per_profile = int(setting(conn, "crawl_per_profile_pages", 5))
        budget = int(setting(conn, "crawl_max_pages", 2000))
    pages = 0
    for pr in profiles:
        if pages >= budget:
            stats["note"] = f"достигнут предел {budget} страниц за прогон"
            break
        with conn_factory() as conn:
            hh = HH(conn, client)
            job = _job(conn, src["id"], pr["id"], hh.search_params(pr))
        try:
            with conn_factory() as conn:
                hh = HH(conn, client)
                items = list(hh.search(pr, min(per_profile, budget - pages)))
            pages += max(1, (len(items) + 49) // 50)
            ingest_items(conn_factory, "hh", iter(items), stats)
            with conn_factory() as conn:
                _finish(conn, job, True)
        except (CrawlBlocked, NotConfigured) as e:
            with conn_factory() as conn:
                _finish(conn, job, False, str(e))
            raise
        except httpx.HTTPError as e:
            with conn_factory() as conn:
                _finish(conn, job, False, str(e))
            stats["errors"] += 1
    return stats


def collect_trudvsem(conn_factory, client: httpx.Client | None = None) -> dict:
    stats = _stats()
    with conn_factory() as conn:
        src = source(conn, "trudvsem")
        if not src["is_active"]:
            return {**stats, "note": "источник выключен"}
        profiles = pipeline.active_profiles(conn)
        pages = min(int(setting(conn, "crawl_per_profile_pages", 5)) * 4, int(setting(conn, "crawl_max_pages", 2000)))
        job = _job(conn, src["id"], None, {"pages": pages})
    try:
        with conn_factory() as conn:
            items = list(Trudvsem(conn, client).cvs(pages))
        ingest_items(conn_factory, "trudvsem", iter(items), stats,
                     filter_fn=lambda n: _relevant(profiles, n.title, n.skills))
        with conn_factory() as conn:
            _finish(conn, job, True)
    except Exception as e:
        with conn_factory() as conn:
            _finish(conn, job, False, str(e))
        raise
    return stats


# ================================================================== выгрузки из кабинетов

def read_table(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t") if text.strip() else csv.excel
        return list(csv.DictReader(io.StringIO(text), dialect=dialect))
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it, [])]
        return [dict(zip(head, r)) for r in it if any(v not in (None, "") for v in r)]
    raise ValueError(f"формат {path.suffix} не поддерживается")


def source_for_file(conn, name: str) -> str:
    """Источник по началу имени файла: habr_*.csv, avito_*.xlsx; остальное — ручной ввод."""
    prefix = name.split("_", 1)[0].lower()
    return prefix if one(conn, "SELECT 1 FROM source WHERE code = %s", (prefix,)) else "manual"


def import_files(conn_factory, folder: str | None = None) -> dict:
    stats = _stats() | {"files": 0}
    root = Path(folder or get_config().import_dir)
    root.mkdir(parents=True, exist_ok=True)
    done_dir = root / "done"
    done_dir.mkdir(exist_ok=True)
    for path in sorted(p for p in root.iterdir() if p.is_file() and p.suffix.lower() in (".csv", ".xlsx", ".xlsm")):
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        with conn_factory() as conn:
            if one(conn, "SELECT 1 FROM exchange_file WHERE kind = 'import' AND sha256 = %s", (sha,)):
                shutil.move(str(path), done_dir / path.name)
                continue
            code = source_for_file(conn, path.name)
        file_stats = _stats()
        try:
            ingest_items(conn_factory, code, iter(read_table(path)), file_stats)
            state, err = "applied", None
        except Exception as e:  # noqa: BLE001
            state, err = "rejected", str(e)
        with conn_factory() as conn:
            run(conn, """INSERT INTO exchange_file (kind, file_name, sha256, state, stats, error)
                         VALUES ('import', %s, %s, %s, %s, %s) ON CONFLICT DO NOTHING""",
                (path.name, sha, state, j(file_stats), err))
        shutil.move(str(path), done_dir / path.name)
        stats["files"] += 1
        for k in ("seen", "new", "changed", "persons", "alerts", "skipped", "errors"):
            stats[k] += file_stats[k]
    return stats


# ================================================================== вакансии с hh

def vacancy_from_hh(conn_factory, client: httpx.Client | None = None) -> dict:
    """Опубликованные компанией вакансии заводятся сами, с origin = hh_publication, вместе с откликами."""
    from .scoring import trigram_sim
    out = {"vacancies": 0, "responses": 0, "unmatched": []}
    with conn_factory() as conn:
        items = HH(conn, client).employer_vacancies()
        profiles = pipeline.active_profiles(conn)
    for v in items:
        with conn_factory() as conn:
            existing = one(conn, "SELECT id FROM vacancy WHERE origin = 'hh_publication' AND external_id = %s", (str(v["id"]),))
            if existing:
                continue
            best = max(profiles, key=lambda p: trigram_sim(v["name"], p["title"]), default=None)
            if not best or trigram_sim(v["name"], best["title"]) < 0.3:
                out["unmatched"].append(v["name"])
                continue
            run(conn, """INSERT INTO vacancy (role_profile_id, title, origin, external_id, opened_at)
                         VALUES (%s, %s, 'hh_publication', %s, %s)""",
                (best["id"], v["name"], str(v["id"]), (v.get("published_at") or "")[:10] or None))
            out["vacancies"] += 1
    with conn_factory() as conn:
        pubs = rows(conn, "SELECT external_id FROM vacancy WHERE origin = 'hh_publication' AND state = 'open'")
    for pub in pubs:
        try:
            with conn_factory() as conn:
                resumes = HH(conn, client).responses(pub["external_id"])
        except (httpx.HTTPError, NotConfigured):
            continue
        s = _stats()
        ingest_items(conn_factory, "hh", iter(resumes), s)
        out["responses"] += s["seen"]
    if out["unmatched"]:
        with conn_factory() as conn:
            from .db import notify_role
            notify_role(conn, "head_recruiter", "vacancy_unmatched",
                        {"text": "Вакансии с hh без подходящего профиля должности", "titles": out["unmatched"]})
    return out
