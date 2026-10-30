"""Выгрузка досье (S10; пункты 25, 30): Word — на сервере, PDF — печать браузера. Обе пишут dossier_export.

Разделы — как в пункте 30 приёмки. Значений контактов в файле нет: только отметка, открыт ли контакт.
Текст из резюме вставляется как текст: python-docx экранирует его сам (W4).
"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document

from .db import j, run, val
from .views import dossier

SECTIONS = ["Сводка", "Сшивка", "Резюме", "Места работы", "Навыки", "Цифровой след", "Публичная активность",
            "Хронология", "Оценки", "Взаимодействие", "Доступ"]


def record(conn, person_id: int, candidate_id: int | None, fmt: str, user_id: int, now: datetime) -> int:
    eid = val(conn, """INSERT INTO dossier_export (person_id, candidate_id, format, author_id, created_at)
                       VALUES (%s, %s, %s, %s, %s) RETURNING id""", (person_id, candidate_id, fmt, user_id, now))
    run(conn, """INSERT INTO audit_log (user_id, action, object_type, object_id, detail, at)
                 VALUES (%s, 'export', 'person', %s, %s, %s)""",
        (user_id, person_id, j({"format": fmt, "id": eid}), now))
    return eid


def _p(doc, text):
    if text:
        doc.add_paragraph(str(text))


def build(conn, person_id: int, now: datetime) -> tuple[bytes, dict] | None:
    d = dossier(conn, person_id, full=True)
    if d is None:
        return None
    card, obs, matches = d["person"], d["observations"], d["matches"]
    doc = Document()
    doc.add_heading(card["name"] or card["label"], 0)
    _p(doc, f"Досье сформировано {now:%d.%m.%Y %H:%M}. Контакты в выгрузку не входят.")
    doc.add_heading(SECTIONS[0], 1)
    best = matches[0] if matches else None
    _p(doc, f"{card['label']}. " + (f"Лучший балл {best['score']} по профилю «{best['role_title']}»." if best else
                                     "Оценок нет."))
    doc.add_heading(SECTIONS[1], 1)
    for o in obs:
        _p(doc, f"{o['source_title']}: {o['url'] or 'ссылки нет'}, видели {o['first_seen']:%d.%m.%Y} — "
                f"{o['last_seen']:%d.%m.%Y}" + (", резюме снято" if o["is_gone"] else ""))
    doc.add_heading(SECTIONS[2], 1)
    for o in obs[:1]:
        _p(doc, f"Должность: {o['title'] or '—'}; город: {o['city'] or '—'}; опыт: {o['exp_years'] or '—'} лет; "
                f"ожидания: {o['salary_expect'] or 'не указаны'}; образование: {o['education'] or '—'}")
        _p(doc, o["about"])
    doc.add_heading(SECTIONS[3], 1)
    for e in (obs[0]["experience"] if obs else []):
        _p(doc, f"{e['company'] or '—'} — {e['position'] or '—'} ({e['start'] or '?'} — {e['end'] or 'по сей день'})")
    doc.add_heading(SECTIONS[4], 1)
    _p(doc, ", ".join(sorted({s for o in obs for s in o["skills"] or []})) or "не указаны")
    doc.add_heading(SECTIONS[5], 1)
    for s in d["social"]:
        _p(doc, f"{s['kind']}: {s['url']} — {s['confidence']}" + (f"; {s['summary']}" if s["summary"] else ""))
    if not d["social"]:
        _p(doc, "Профилей в сети не найдено или не подтверждено.")
    doc.add_heading(SECTIONS[6], 1)
    acts = [s for s in d["social"] if s.get("topics") or s.get("last_activity")]
    for s in acts:
        _p(doc, f"{s['kind']}: темы {', '.join(s['topics'] or []) or '—'}, последняя активность "
                f"{s['last_activity'] or '—'}")
    if not acts:
        _p(doc, "Публичной профессиональной активности не видно.")
    doc.add_heading(SECTIONS[7], 1)
    for e in d["events"][:40]:
        _p(doc, f"{e['happened_at']:%d.%m.%Y}: {e['kind']}")
    doc.add_heading(SECTIONS[8], 1)
    for m in matches:
        _p(doc, f"«{m['role_title']}» (версия {m['profile_version']}): {m['score']} из 100")
        _p(doc, m["explanation"])
    doc.add_heading(SECTIONS[9], 1)
    c = d["candidate"]
    _p(doc, f"Кандидат: статус {c['status']}, в работе с {c['created_at']:%d.%m.%Y}; контакт "
            f"{'открыт' if c['contact_opened'] else 'не открыт'}." if c else "В работу не брали.")
    for x in d["decisions"]:
        _p(doc, f"{x['created_at']:%d.%m.%Y}, {x['author']}: {x['kind']} — {x['comment']}")
    doc.add_heading(SECTIONS[10], 1)
    _p(doc, "Просмотры карточки руководителями, открытия контактов и выгрузки пишутся в журнал доступа.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), d
