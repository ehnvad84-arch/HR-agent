# -*- coding: utf-8 -*-
"""Документы проекта из Markdown в оформленный .docx.

    python build_tz_docx.py  →  ТЗ_для_инженеров_развёртывание.docx и Мануал_рекрутера.docx

Источник один — Markdown. Документы Word руками не правятся: правится Markdown и
пересобирается, иначе две версии разъедутся.
Поддерживается: заголовки, абзацы, списки, таблицы, блоки кода, **жирный** и `код` в строке.
"""
import pathlib
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = [
    ("ТЗ_для_инженеров_развёртывание.md", "ТЗ для инженеров"),
    ("Мануал_рекрутера.md", "Мануал рекрутера"),
]

INK = RGBColor(0x1A, 0x1A, 0x1A)
GREEN = RGBColor(0x1E, 0x7B, 0x45)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
CODE_BG = "F2F4F3"
HEAD_BG = "E6F4EA"


def shade(cell_or_par, color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), color)
    target = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else \
        cell_or_par.paragraph_format.element.get_or_add_pPr()
    target.append(el)


def rich(par, text, size=11, color=INK):
    """**жирный** и `моноширинный` внутри строки."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(size - 1)
            r.font.color.rgb = GREEN
        else:
            r = par.add_run(part)
        if not r.font.color.rgb:
            r.font.color.rgb = color
        if not r.font.name:
            r.font.name = "Arial"
        if not r.font.size:
            r.font.size = Pt(size)


def heading(doc, text, level):
    p = doc.add_heading("", level=min(level, 4))
    for run in p.runs:
        run.clear()
    r = p.add_run(re.sub(r"[`*]", "", text))
    r.font.name = "Arial"
    r.font.color.rgb = GREEN if level <= 2 else INK
    r.font.size = Pt({1: 18, 2: 15, 3: 13, 4: 12}.get(level, 12))
    r.bold = True
    return p


def code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    shade(p, CODE_BG)
    r = p.add_run("\n".join(lines))
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = INK


def table_block(doc, rows):
    cols = len(rows[0])
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row[:cols]):
            cells[j].text = ""
            par = cells[j].paragraphs[0]
            rich(par, val, size=9.5)
            if i == 0:
                for r in par.runs:
                    r.bold = True
                shade(cells[j], HEAD_BG)
    doc.add_paragraph()
    return t


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build(src_name, podpis):
    SRC = pathlib.Path(src_name)
    OUT = SRC.with_suffix(".docx")
    if not SRC.exists():
        sys.exit(f"нет исходника {SRC}")
    md = SRC.read_text(encoding="utf-8").splitlines()

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(1.8)
        foot = s.footer.paragraphs[0]
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = foot.add_run(f"Агент подбора персонала · {podpis} · внутренний документ")
        fr.font.size = Pt(8)
        fr.font.color.rgb = GREY

    i, in_code, buf = 0, False, []
    counts = {"h": 0, "p": 0, "t": 0, "code": 0, "li": 0}
    while i < len(md):
        line = md[i]

        if line.strip().startswith("```"):
            if in_code:
                code_block(doc, buf); counts["code"] += 1; buf, in_code = [], False
            else:
                in_code = True
            i += 1; continue
        if in_code:
            buf.append(line); i += 1; continue

        if not line.strip():
            i += 1; continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            heading(doc, m.group(2), len(m.group(1))); counts["h"] += 1; i += 1; continue

        if line.strip().startswith("|") and i + 1 < len(md) and re.match(r"^\s*\|[\s:|-]+\|\s*$", md[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < len(md) and md[i].strip().startswith("|"):
                rows.append(split_row(md[i])); i += 1
            table_block(doc, rows); counts["t"] += 1; continue

        if line.strip() in ("---", "***"):
            sep = doc.add_paragraph()
            sep.paragraph_format.space_before = Pt(2)
            sep.paragraph_format.space_after = Pt(2)
            i += 1; continue

        m = re.match(r"^\s*([-*]|\d+\.)\s+(.*)$", line)
        if m:
            style = "List Number" if m.group(1)[0].isdigit() else "List Bullet"
            chunk = [m.group(2)]
            i += 1
            while i < len(md) and md[i].startswith("   ") and md[i].strip() and not re.match(
                    r"^\s*([-*]\s|\d+\.\s)", md[i]):
                chunk.append(md[i].strip()); i += 1
            p = doc.add_paragraph(style=style)
            rich(p, " ".join(chunk))
            counts["li"] += 1; continue

        if line.strip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            rich(p, line.strip().lstrip("> "), color=GREY)
            counts["p"] += 1; i += 1; continue

        # абзац собирается целиком: в markdown он разбит переносами,
        # иначе жирный, начатый на одной строке и закрытый на другой, остаётся сырым
        chunk = [line.strip()]
        i += 1
        while i < len(md) and md[i].strip() and not re.match(
                r"^\s*(#{1,4}\s|[-*]\s|\d+\.\s|\||>|```|---$|\*\*\*$)", md[i]):
            chunk.append(md[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        rich(p, " ".join(chunk))
        counts["p"] += 1

    doc.save(OUT)
    print(f"собрано {OUT}: заголовков {counts['h']}, абзацев {counts['p']}, "
          f"таблиц {counts['t']}, блоков кода {counts['code']}, пунктов списков {counts['li']}")
    # проверка: документ читается обратно и не потерял содержимое
    back = Document(OUT)
    assert len(back.paragraphs) > 80, f"{OUT}: подозрительно мало абзацев"
    assert len(back.tables) == counts["t"], f"{OUT}: таблиц в файле меньше, чем собрано"
    print(f"  проверка: в файле {len(back.paragraphs)} абзацев и {len(back.tables)} таблиц")


if __name__ == "__main__":
    for name, podpis in DOCS:
        build(name, podpis)
