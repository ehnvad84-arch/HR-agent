"""Собирает шаблон выгрузки с Хабр Карьеры: docs/templates/habr_шаблон.xlsx.

Колонки совпадают с FILE_ALIASES в server/hr/normalize.py. Запуск: python docs/templates/make_habr_template.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

COLUMNS = [
    ("Ссылка", "Ссылка на профиль на career.habr.com. Обязательно: по ней агент узнаёт человека при повторной выгрузке", 45),
    ("Должность", "Желаемая или текущая должность", 30),
    ("Город", "Город проживания", 16),
    ("Навыки", "Через запятую: Altium, C, STM32", 40),
    ("Места работы", "Компании через запятую, текущая первой", 40),
    ("Текущий работодатель", "Если известен отдельно", 25),
    ("Опыт лет", "Число, например 6 или 6,5", 10),
    ("Ожидания", "Зарплата в рублях, только число", 12),
    ("Переезд", "да / нет", 9),
    ("Ищет работу", "да / нет — статус на Хабре", 12),
    ("Обновлено", "Дата обновления профиля, ДД.ММ.ГГГГ", 12),
    ("Образование", "Вуз и специальность", 30),
    ("О себе", "Коротко, можно пусто", 40),
    ("Ник", "Ник на Хабре или GitHub — помогает сшить с другими источниками", 16),
    ("ФИО", "Если видно в профиле", 25),
]

wb = Workbook()
ws = wb.active
ws.title = "Резюме"
for i, (name, hint, width) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    c.comment = Comment(hint, "агент подбора")
    ws.column_dimensions[c.column_letter].width = width
ws.freeze_panes = "A2"
wb.save(Path(__file__).with_name("habr_шаблон.xlsx"))
